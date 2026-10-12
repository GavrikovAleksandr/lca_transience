"""Print the portable Morocco PV snapshot; never import or mutate Brightway.

Run with the bundled Python runtime. The source SQLite database is opened with
mode=ro in one explicit read transaction. Only the selected reference activity,
its required auxiliary activities, and background activity identities are read.
Output is JSON on stdout; save it using the task's normal apply_patch workflow.
"""

from __future__ import annotations

import argparse
from collections import deque
from copy import deepcopy
import hashlib
import json
import math
from pathlib import Path
import pickle
import sqlite3

import openpyxl


PROVIDER = "electricity production, photovoltaic, Morocco (por)"
PRODUCT = "electricity, medium voltage"
SOURCE_NAME = "electricity production, photovoltaic, commercial"
DEFAULT_PROJECT = Path(
    "C:/Users/terlouw_t/AppData/Local/pylca/Brightway3/por.7cf41ff9/lci/databases.db"
)
DEFAULT_WORKBOOK = Path(
    "C:/Users/terlouw_t/.conda/envs/premise_pathways/Lib/site-packages/"
    "premise/data/additional_inventories/lci-PV.xlsx"
)
REFERENCE = "ecoinvent_312_reference"
BACKGROUND = "ecoinvent-3.12-cutoff"
BIOSPHERE = "ecoinvent-3.12-biosphere"


def identity(record):
    return (
        record.get("name"),
        record.get("product", record.get("reference product")),
        record.get("location"),
        record.get("unit"),
    )


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def workbook_source(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parameters = list(workbook["parameters"].iter_rows(values_only=True))
    lifetime = parameters[180][2]
    yields = {"roof": parameters[195][14], "facade": parameters[196][14],
              "ground": parameters[197][14]}
    assert parameters[194][14] == "MA"
    assert lifetime == 30
    assert math.isclose(yields["roof"], 0.94 * yields["ground"], rel_tol=1e-14)
    assert math.isclose(yields["facade"], 0.66 * yields["ground"], rel_tol=1e-14)
    rows = list(workbook["PV (extension)"].iter_rows(
        min_row=1155, max_row=1180, values_only=True
    ))
    assert rows[0][0:2] == ("Activity", SOURCE_NAME)
    assert rows[2][0:2] == ("location", "MA")
    headers = rows[9]
    exchanges = [dict((key, value) for key, value in zip(headers, row)
                      if key is not None and value is not None) for row in rows[10:]]
    workbook.close()
    return lifetime, yields, exchanges


def canonical(record):
    result = deepcopy(record)
    for field in ("database", "code", "input", "output", "worksheet name",
                  "id", "_id", "_document_id", "filename"):
        result.pop(field, None)
    if "product" in result:
        result["reference product"] = result.pop("product")
    if result.get("type") == "processwithreferenceproduct":
        result["type"] = "process"
    return result


def build(project, workbook):
    lifetime, yields, original_exchanges = workbook_source(workbook)
    connection = sqlite3.connect(project.resolve().as_uri() + "?mode=ro", uri=True)
    connection.execute("PRAGMA query_only = ON")
    connection.execute("BEGIN")
    try:
        background = {}
        for code, name, product, location, data in connection.execute(
                'SELECT code, name, product, location, data FROM activitydataset '
                'WHERE "database" = ?', (BACKGROUND,)):
            metadata = pickle.loads(data)
            key = (name, product, location, metadata["unit"])
            background.setdefault(key, []).append(code)
        if not background:
            raise ValueError("The named ecoinvent 3.12 cutoff database is empty")

        metadata_cache = {}
        def metadata(database, code):
            key = (database, code)
            if key not in metadata_cache:
                matches = connection.execute(
                    'SELECT data FROM activitydataset WHERE "database" = ? AND code = ?',
                    key,
                ).fetchall()
                if len(matches) != 1:
                    raise ValueError(f"Missing or ambiguous source activity: {key}")
                metadata_cache[key] = pickle.loads(matches[0][0])
            return metadata_cache[key]

        def activity(database, code):
            result = canonical(metadata(database, code))
            result["exchanges"] = []
            for (data,) in connection.execute(
                    'SELECT data FROM exchangedataset WHERE output_database = ? '
                    'AND output_code = ? ORDER BY id', (database, code)):
                original = pickle.loads(data)
                exchange = canonical(original)
                if exchange["type"] == "technosphere":
                    supplier = metadata(*original["input"])
                    for field in ("name", "reference product", "location", "unit"):
                        exchange[field] = supplier[field]
                elif exchange["type"] == "biosphere":
                    supplier = metadata(*original["input"])
                    if original["input"][0] != BIOSPHERE:
                        raise ValueError("Unexpected biosphere database")
                    exchange["name"] = supplier["name"]
                    exchange["unit"] = supplier["unit"]
                    exchange["categories"] = list(supplier["categories"])
                result["exchanges"].append(exchange)
            return result

        matches = connection.execute(
            'SELECT code FROM activitydataset WHERE "database" = ? AND name = ? '
            'AND product = ? AND location = ?',
            (REFERENCE, SOURCE_NAME, "electricity, low voltage", "MA"),
        ).fetchall()
        if len(matches) != 1:
            raise ValueError("Missing or ambiguous migrated Morocco PV source")
        source_code = matches[0][0]
        source_metadata_sha256 = hashlib.sha256(json.dumps(
            canonical(metadata(REFERENCE, source_code)), sort_keys=True,
            ensure_ascii=True, separators=(",", ":"), allow_nan=False,
        ).encode("utf-8")).hexdigest()
        source = activity(REFERENCE, source_code)
        source_recipe_sha256 = hashlib.sha256(json.dumps(
            source, sort_keys=True, ensure_ascii=True, separators=(",", ":"),
            allow_nan=False,
        ).encode("utf-8")).hexdigest()
        raw = {e["name"]: e["amount"] for e in original_exchanges}
        for exchange in source["exchanges"]:
            if exchange["name"] in raw:
                if not math.isclose(exchange["amount"], raw[exchange["name"]],
                                    rel_tol=1e-14, abs_tol=0):
                    raise ValueError("Source workbook and migrated coefficients differ")
        source["name"] = PROVIDER
        source["reference product"] = PRODUCT
        source["type"] = "process"
        source["pv lifetime years"] = lifetime
        source["pv annual yield kwh per kwp"] = yields
        source["pv mix origin"] = "Existing Morocco commercial adapter: source-country generation-weighted proxy"
        source["pv yield origin"] = "Global Solar Atlas numeric country PVOUT retained in lci-PV.xlsx"
        source["comment"] += (
            "\nPoR compatibility snapshot of the existing Morocco commercial PV recipe, "
            "migrated to ecoinvent 3.12. The medium-voltage output is a lossless plant "
            "interface to the original low-voltage-labelled recipe: no unreported "
            "transformer, losses, distribution service or grid generation is added. "
            "The source proxy technology shares, lifetime, cleaning quantities "
            "and manufacturing/construction supplier geographies are preserved. "
            "Country PVOUT retrieval date was not recorded in the source workbook. "
            "The source cleaning quantities are retained as supplied; they are not "
            "recomputed from module area or Morocco yield."
            " The source calls this an OECD proxy, but its formulas weight all 33 "
            "report countries, including some countries outside the OECD."
        )
        for exchange in source["exchanges"]:
            if exchange["type"] == "production":
                exchange.update(name=PROVIDER, **{"reference product": PRODUCT})
                exchange.pop("production volume", None)

        pending = deque()
        for (data,) in connection.execute(
                'SELECT data FROM exchangedataset WHERE output_database = ? AND '
                'output_code = ? AND type = ?', (REFERENCE, source_code, "technosphere")):
            exchange = pickle.loads(data)
            supplier = metadata(*exchange["input"])
            if identity(supplier) not in background:
                pending.append(tuple(exchange["input"]))
        auxiliaries = {}
        while pending:
            key = pending.popleft()
            if key in auxiliaries:
                continue
            if key[0] != REFERENCE:
                raise ValueError(f"Unexpected auxiliary database: {key[0]}")
            auxiliaries[key] = activity(*key)
            for (data,) in connection.execute(
                    'SELECT data FROM exchangedataset WHERE output_database = ? AND '
                    'output_code = ? AND type = ?', (*key, "technosphere")):
                exchange = pickle.loads(data)
                supplier = metadata(*exchange["input"])
                if identity(supplier) not in background:
                    pending.append(tuple(exchange["input"]))

        ordered = sorted(auxiliaries.values(), key=identity)
        datasets = [source] + ordered
        bundled = {identity(dataset) for dataset in datasets}
        external = set()
        biosphere = set()
        for dataset in datasets:
            productions = [e for e in dataset["exchanges"] if e["type"] == "production"]
            if len(productions) != 1 or identity(productions[0]) != identity(dataset):
                raise ValueError(f"Invalid production identity: {identity(dataset)}")
            for exchange in dataset["exchanges"]:
                if exchange["type"] == "technosphere":
                    key = identity(exchange)
                    if key not in bundled:
                        if key not in background or len(background[key]) != 1:
                            raise ValueError(f"Missing or ambiguous background supplier: {key}")
                        external.add(key)
                elif exchange["type"] == "biosphere":
                    biosphere.add((exchange["name"], tuple(exchange["categories"]), exchange["unit"]))
        installations = [e for e in source["exchanges"] if e.get("scope") == "commercial"]
        if len(installations) != 11:
            raise ValueError("Expected all eleven original Morocco PV technology inputs")
        normalised = sum(e["amount"] * e["kwp"] * e["productivity"] * e["lifetime"]
                         for e in installations)
        # Preserve the existing coefficients, including their small source
        # normalisation discrepancy. Native PV mix validation permits 0.005.
        if not math.isclose(normalised, 1, abs_tol=0.005):
            raise ValueError(f"Installation electricity shares sum to {normalised}")
        result = {
            "schema_version": 1,
            "ecoinvent_version": "3.12",
            "system_model": "cutoff",
            "provider": {field: source[field] for field in ("name", "reference product", "unit", "location")},
            "provenance": {
                "source_description": "Existing Morocco commercial PV adapter and required legacy PV auxiliaries, snapshotted from an already migrated ecoinvent 3.12 reference database",
                "source_path": "premise/data/additional_inventories/lci-PV.xlsx",
                "source_sha256": sha256(workbook),
                "source_workbook": "premise 2.3.8 data/additional_inventories/lci-PV.xlsx",
                "source_workbook_sha256": sha256(workbook),
                "source_recipe_sha256": source_recipe_sha256,
                "source_activity_metadata_sha256": source_metadata_sha256,
                "source_inventory_version": "3.7 (native premise import declaration)",
                "snapshot_inventory_version": "3.12 (existing migrated reference record)",
                "source_record": {"sheet": "PV (extension)", "range": "A1155:J1180"},
                "source_parameters": {"sheet": "parameters", "lifetime_cell": "C181",
                                      "ground_yield_cell": "O198", "roof_yield_cell": "O196",
                                      "facade_yield_cell": "O197", "country_cell": "O195"},
                "source_atlas": "https://globalsolaratlas.info/download/world",
                "source_atlas_numeric_cell": {"sheet": "extension - productivity", "range": "C20"},
                "source_atlas_retrieval_date": "not recorded in source workbook",
                "source_pv_lci": "https://doi.org/10.13140/RG.2.2.17977.19041",
                "mix_assumption": "Existing Morocco adapter: generation-weighted technology proxy from all 33 source report countries, normalised within commercial segment",
                "source_mix_description_caveat": "Workbook prose describes an OECD-country proxy, but parameter formulas include all 33 report countries in rows 6:38, including CN, MY, ZA and TH",
                "lifetime_years": lifetime,
                "annual_yield_kwh_per_kwp": yields,
                "yield_factors": {"roof": 0.94, "facade": 0.66, "ground": 1},
                "scope_share": installations[0]["scope share"],
                "source_coefficient_equality": "Every unchanged-name Morocco exchange amount equals the original cached workbook value",
                "output_interface": "Original LV recipe is exposed at MV as a lossless plant interface; no additional transformer or grid mix",
                "cleaning_assumption": "Original tap water 0.0037026 kg/kWh and wastewater -3.3323e-06 m3/kWh preserved",
                "source_production_volume": {"value": 27.6, "note": "Source volume retained only as provenance; excluded from new production exchange because source comment year/unit are ambiguous"},
                "auxiliary_count": len(ordered),
                "background_identity_count": len(background),
                "external_supplier_count": len(external),
                "biosphere_flow_count": len(biosphere),
                "installation_electricity_share_sum": normalised,
                "source_share_normalisation_discrepancy": normalised - 1,
                "source_share_normalisation_policy": "Original coefficients retained; inferred electricity shares differ from one by 0.03653010668472%, within native premise PV tolerance 0.005",
                "validation": "All technosphere links resolve uniquely to bundled auxiliaries or unchanged ecoinvent 3.12 cutoff identities; all biosphere identities copied from linked ecoinvent 3.12 biosphere records",
            },
            "datasets": datasets,
        }
        return result
    finally:
        connection.rollback()
        connection.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-db", type=Path, default=DEFAULT_PROJECT)
    parser.add_argument("--source-workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--summary", action="store_true")
    parser.add_argument("--compact", action="store_true")
    parser.add_argument("--length", action="store_true")
    parser.add_argument("--chunk-offset", type=int)
    parser.add_argument("--chunk-size", type=int, default=20000)
    args = parser.parse_args()
    result = build(args.project_db, args.source_workbook)
    if args.summary:
        result = {"provider": result["provider"], "provenance": result["provenance"],
                  "auxiliaries": [list(identity(d)) for d in result["datasets"][1:]]}
    encoded = json.dumps(result, indent=None if args.compact else 2,
                         ensure_ascii=True, allow_nan=False,
                         separators=(",", ":") if args.compact else None)
    if args.length:
        print(len(encoded))
    elif args.chunk_offset is not None:
        print(encoded[args.chunk_offset:args.chunk_offset + args.chunk_size])
    else:
        print(encoded)


if __name__ == "__main__":
    main()
