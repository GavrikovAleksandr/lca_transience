"""Optional read-only workbook-to-YAML extraction; no work runs on import."""
from __future__ import annotations

import argparse
from pathlib import Path
import os
import tempfile
import yaml

if __package__:
    from .share_files import DATA_DIR, ShareDataError, _safe_output, validate_share_mapping
else:
    from share_files import DATA_DIR, ShareDataError, _safe_output, validate_share_mapping


def read_excel_shares(input_xlsx=DATA_DIR / "metal_primary_secondary_shares.xlsx"):
    from openpyxl import load_workbook
    workbook = load_workbook(Path(input_xlsx).resolve(), read_only=True, data_only=True)
    try:
        if "Shares" not in workbook.sheetnames:
            raise ShareDataError(f"{input_xlsx}: required worksheet 'Shares' is missing")
        result, blank_seen = {}, False
        for row_number, row in enumerate(workbook["Shares"].iter_rows(min_row=5, max_col=7, values_only=True), 5):
            key, name, reference, p2020, p2050, s2020, s2050 = row
            if key in (None, ""):
                if any(value not in (None, "") for value in row):
                    raise ShareDataError(f"{input_xlsx}: row {row_number} has data but no product key")
                blank_seen = True
                continue
            if blank_seen:
                raise ShareDataError(f"{input_xlsx}: product rows are interrupted before row {row_number}; do not silently truncate the workbook")
            key = str(key)
            if key in result:
                raise ShareDataError(f"{input_xlsx}: duplicate product {key!r} at row {row_number}")
            result[key] = {"name": name, "reference product": reference,
                           "shares": {"primary": {2020: p2020, 2050: p2050},
                                      "secondary": {2020: s2020, 2050: s2050}}}
        return validate_share_mapping(result, label=str(input_xlsx))
    finally:
        workbook.close()


def convert_workbook_to_yaml(input_xlsx=DATA_DIR / "metal_primary_secondary_shares.xlsx",
                             output_yaml=DATA_DIR / "metal_primary_secondary_shares.yaml", *, force=False):
    result = read_excel_shares(input_xlsx)
    target = _safe_output(output_yaml)
    payload = yaml.safe_dump(result, sort_keys=False, allow_unicode=True).encode("utf-8")
    if target.is_file() and target.read_bytes() == payload:
        return {"status": "reused", "product_count": len(result), "output_yaml": str(target)}
    if target.exists() and not force:
        raise ShareDataError(f"{target} already exists. Use the YAML source directly, or explicitly force workbook extraction after reviewing it.")
    target.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=".metal-baseline-", suffix=".tmp", dir=target.parent)
    try:
        with os.fdopen(descriptor, "wb") as stream:
            stream.write(payload)
        os.replace(temporary, target)
    finally:
        if Path(temporary).exists():
            Path(temporary).unlink()
    return {"status": "extracted", "product_count": len(result), "output_yaml": str(target)}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, allow_abbrev=False)
    parser.add_argument("--input-xlsx", type=Path, default=DATA_DIR / "metal_primary_secondary_shares.xlsx")
    parser.add_argument("--output-yaml", type=Path, default=DATA_DIR / "metal_primary_secondary_shares.yaml")
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args(argv)
    summary = convert_workbook_to_yaml(**vars(args))
    print(f"{summary['status']}: {summary['product_count']} products, {summary['output_yaml']}")
    return summary


if __name__ == "__main__":
    main()
